#!/usr/bin/env python3
"""批量入库：把目录/文件里的文档喂给知识库。

用法:
  python scripts/ingest.py                # 入库 docs/ 目录（默认）
  python scripts/ingest.py docs 报告      # 入库多个目录/文件
  python scripts/ingest.py --file a.md    # 入库单个文件
  python scripts/ingest.py --clean        # 清空知识库（危险，需确认）

支持格式: .md .markdown .txt .pdf .docx
"""
import argparse
import asyncio
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from app.rag_factory import create_rag, llm_binding_info, embedding_binding_info, storage_info  # noqa: E402

SUPPORTED = {".md", ".markdown", ".txt", ".pdf", ".docx"}


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # 表格内容也抽出来
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


READERS = {".pdf": read_pdf, ".docx": read_docx}

CATEGORIES = ["产品资产", "行业知识", "服务知识", "创作成果"]


def write_categories_manifest(files: list[Path]) -> None:
    """把 文件名→分类 映射写到 WebUI 静态目录，资产库视图按它恢复分类。"""
    import json

    mapping: dict[str, str] = {}
    for f in files:
        parts = list(f.parts)
        if "docs" in parts:
            i = parts.index("docs")
            if len(parts) > i + 1 and parts[i + 1] in CATEGORIES:
                mapping[f.name] = parts[i + 1]
    out = PROJECT_ROOT / "app" / "webui_cn" / "categories.json"
    out.write_text(
        json.dumps(mapping, ensure_ascii=False, indent=0), encoding="utf-8"
    )
    print(f"分类清单已更新: {len(mapping)} 个文件 -> {out.name}")


def collect_files(paths: list[str]) -> list[Path]:
    files: list[Path] = []
    for p in paths:
        path = Path(p)
        if not path.exists():
            print(f"⚠️  跳过不存在的路径: {p}")
            continue
        if path.is_file():
            if path.suffix.lower() in SUPPORTED:
                files.append(path)
            else:
                print(f"⚠️  跳过不支持的格式: {p}（支持 {', '.join(sorted(SUPPORTED))}）")
        else:
            files.extend(
                f for f in sorted(path.rglob("*"))
                if f.is_file() and f.suffix.lower() in SUPPORTED
                and not f.name.startswith((".", "~"))
            )
    return files


async def cmd_clean() -> None:
    wd = Path(storage_info()["working_dir"])
    if not wd.exists():
        print("存储目录不存在，无需清理")
        return
    answer = input(f"⚠️  将删除 {wd} 下全部知识库数据（向量/图谱/缓存），输入 yes 确认: ")
    if answer.strip().lower() != "yes":
        print("已取消")
        return
    import shutil

    shutil.rmtree(wd)
    print("✅ 已清空，下次入库会自动重建")


async def cmd_ingest(paths: list[str]) -> None:
    files = collect_files(paths)
    if not files:
        print("没有找到可入库的文档。把 .md/.txt/.pdf/.docx 放进 docs/ 目录后重试。")
        return

    # 生成「文件名 → 业务分类」映射清单，供 WebUI 资产库按分类展示
    # （LightRAG 会把 file_path 规范化为纯文件名，目录前缀会丢，所以单独落一份映射）
    write_categories_manifest(files)

    print("模型配置:")
    llm, emb = llm_binding_info(), embedding_binding_info()
    print(f"  LLM       : {llm['binding']} / {llm['model']}  ({llm['host']})")
    print(f"  Embedding : {emb['binding']} / {emb['model']} dim={emb['dim']}")
    sto = storage_info()
    print(f"  向量库    : {sto['vector_storage']} @ {sto['milvus_uri']}")
    print(f"共 {len(files)} 个文件待入库\n")

    rag = await create_rag()
    ok, fail = 0, 0
    for i, f in enumerate(files, 1):
        rel = f.relative_to(PROJECT_ROOT) if f.is_relative_to(PROJECT_ROOT) else f
        try:
            reader = READERS.get(f.suffix.lower(), read_text)
            content = reader(f).strip()
            if not content:
                print(f"[{i}/{len(files)}] ⚠️  空文档，跳过: {rel}")
                continue
            print(f"[{i}/{len(files)}] 入库中: {rel} ({len(content)} 字符)")
            await rag.ainsert(content, file_paths=str(rel))
            ok += 1
        except KeyboardInterrupt:
            print("\n中断：已完成的部分保留")
            break
        except Exception as e:  # 单个文件失败不影响其余
            fail += 1
            print(f"[{i}/{len(files)}] ❌ 失败: {rel} -> {e}")

    try:
        status = await rag.get_docs_by_status()
        counts = {k: len(v) for k, v in status.items()}
        print(f"\n✅ 完成：成功 {ok}，失败 {fail}。库内文档状态: {counts}")
    except Exception:
        print(f"\n✅ 完成：成功 {ok}，失败 {fail}")

    await rag.finalize_storages()


def main() -> None:
    parser = argparse.ArgumentParser(description="RAG 知识库批量入库")
    parser.add_argument("paths", nargs="*", default=["docs"], help="目录或文件（默认 docs/）")
    parser.add_argument("--file", dest="single", help="入库单个文件")
    parser.add_argument("--clean", action="store_true", help="清空知识库")
    args = parser.parse_args()

    if args.clean:
        asyncio.run(cmd_clean())
    elif args.single:
        asyncio.run(cmd_ingest([args.single]))
    else:
        asyncio.run(cmd_ingest(args.paths))


if __name__ == "__main__":
    main()
